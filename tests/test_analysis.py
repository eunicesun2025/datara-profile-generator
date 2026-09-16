import io
import json
import time
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image

from datara.app import create_app
from datara.domain import FieldDef, TableDef, FIELD_NAME, new_profile, normalize, infer_type, validate_profile
from datara.generators import export_zip, prompt
from datara.importer import parse_fields
from datara.provider import parse_analysis, draft_prompt
from datara.references import reference_text


@pytest.mark.parametrize("name,description,expected", [
    ("InvoiceDate", "", "Date"), ("amount", "", "Decimal"), ("unit_price", "", "Decimal"),
    ("quantity", "", "Decimal"), ("account_no", "银行账号", "String"),
    ("invoice_number", "", "String"), ("value", "发票日期", "Date"), ("page_count", "", "Integer"),
])
def test_semantic_type_defaults(name, description, expected):
    assert infer_type(name, description) == expected


def test_semantic_type_overrides_wrong_model_proposal():
    assert infer_type("payment_date", "付款日期", "Decimal") == "Date"
    assert infer_type("total_amount", "总金额", "Integer") == "Decimal"
    assert infer_type("account_no", "账号", "Decimal") == "String"


def test_normalization_is_unique_stable_and_agrees_across_artifacts():
    p = new_profile()
    p.tables[0].fields += [FieldDef(name="invoice_date"), FieldDef(name="InvoiceDate"),
                           FieldDef(name="开户机构"), FieldDef(name="Invoice  Date"), FieldDef(name="123 Code")]
    notes = normalize(p)
    names = [f.name for f in p.tables[0].fields]
    assert len(set(names)) == len(names)
    assert all(FIELD_NAME.fullmatch(name) for name in names)
    assert "invoice_date_2" in names and "invoice_date_3" in names
    assert notes and not normalize(p)
    archive = zipfile.ZipFile(io.BytesIO(export_zip(p)))
    rows = list(load_workbook(io.BytesIO(archive.read("field_mapping.xlsx"))).active.values)[1:]
    structure = json.loads(archive.read("output_structure.json"))
    assert {r[3] for r in rows if r[8] == "AI"} == set(structure["AI_Document"])
    for name in structure["AI_Document"]:
        assert f"[{name}]" in archive.read("create_tables.sql").decode()


def test_reserved_name_collision_is_not_silently_reclassified():
    p = new_profile()
    p.tables[0].fields.append(FieldDef(name="FileID"))
    with pytest.raises(ValueError, match="System"):
        normalize(p)


def test_import_semantics_and_display_are_exported():
    wb = Workbook(); ws = wb.active
    ws.append(["字段", "说明"])
    ws.append(["InvoiceDate", "发票日期"])
    ws.append(["amount", "金额"])
    ws.append(["account_no", "账号"])
    data = io.BytesIO(); wb.save(data)
    p, _ = parse_fields(data.getvalue(), ws.title, {"name": "字段", "description": "说明"})
    fields = [f for f in p.tables[0].fields if f.source == "AI"]
    assert [(f.name, f.data_type, f.head_display) for f in fields] == [
        ("invoice_date", "Date", 1), ("amount", "Decimal", 2), ("account_no", "String", 3)]
    assert not validate_profile(p)["errors"]


def test_analysis_proposes_existing_field_updates_without_mutating_profile():
    p = new_profile()
    p.tables[0].fields.append(FieldDef(name="amount", sql_type="decimal(23,2)", head_display=2))
    before = p.model_dump()
    raw = json.dumps({"profile": {"accepted_documents": "收款凭证", "document_rules": "结合标题及收款标签判断。"},
                      "fields": [{"table_name": "AI_Document", "name": "amount", "data_type": "String",
                                  "description": "金额", "extraction": "读取合计右侧金额。", "head_display": 1},
                                 {"table_name": "AI_Document", "name": "PaymentDate", "data_type": "String"},
                                 {"table_name": "AI_Document", "name": "AccountNo", "data_type": "Decimal"},
                                 {"table_name": "AI_Document", "name": "FileID"}]})
    result = parse_analysis(raw, p)
    assert p.model_dump() == before
    amount, date, account = [s["field"] for s in result["suggestions"]]
    assert result["suggestions"][0]["action"] == "update"
    assert amount["id"] == p.tables[0].fields[-1].id
    assert amount["data_type"] == "Decimal" and amount["sql_type"] == "decimal(23,2)"
    assert date["name"] == "payment_date" and date["data_type"] == "Date"
    assert account["data_type"] == "String"
    assert len({f["head_display"] for f in (amount, date, account)}) == 3
    assert result["profile_suggestion"]["accepted_documents"] == "收款凭证"


def test_analysis_preserves_explicit_head_display_before_filling_gaps():
    p = new_profile()
    result = parse_analysis(json.dumps({"fields": [
        {"table_name": "AI_Document", "name": "quantity"},
        {"table_name": "AI_Document", "name": "customer_name", "head_display": 2},
        {"table_name": "AI_Document", "name": "billing_date"},
        {"table_name": "AI_Document", "name": "account_no", "head_display": 1},
        {"table_name": "AI_Document", "name": "amount", "head_display": 4},
    ]}), p)
    displays = {s["field"]["name"]: s["field"]["head_display"] for s in result["suggestions"]}
    assert displays["account_no"] == 1
    assert displays["customer_name"] == 2
    assert displays["amount"] == 4
    assert displays["quantity"] == 3
    assert displays["billing_date"] == 5


def test_analysis_respects_manual_system_and_detail_display():
    p = new_profile()
    p.tables[0].fields.append(FieldDef(name="note", source="Manual"))
    p.tables.append(TableDef(name="AI_Invoice_Detail", role="detail", parent_table_id=p.tables[0].id))
    normalize(p)
    result = parse_analysis(json.dumps({"fields": [
        {"table_name": "AI_Document", "name": "note"},
        {"table_name": "AI_Invoice_Detail", "name": "item"},
        {"table_name": "AI_Invoice_Detail", "name": "LineAmount", "head_display": 1},
    ]}), p)
    assert len(result["suggestions"]) == 1
    assert result["suggestions"][0]["field"]["head_display"] is None


def test_draft_context_and_final_prompt_separate_reference_and_sample_data():
    p = new_profile()
    p.tables[0].fields.append(FieldDef(name="amount", description="收款金额", data_type="Decimal",
                                     sample_data="DO_NOT_COPY_THIS", extraction="读取总计右侧数字。"))
    draft = draft_prompt(p, "请检查")
    assert "收款金额" in draft and "head_display" in draft and "document_rules" in draft
    extraction = prompt(p)
    assert "读取总计右侧数字" in extraction and "千分位" in extraction
    assert "DO_NOT_COPY_THIS" not in extraction and "file_id" not in extraction


def test_reference_formats_and_limits():
    assert reference_text("字段说明".encode(), "说明.txt") == "字段说明"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Invoice rules</w:t></w:r></w:p></w:body></w:document>')
    assert reference_text(out.getvalue(), "rules.docx") == "Invoice rules"
    with pytest.raises(ValueError, match="60,000"):
        reference_text(b"a" * 60001, "large.txt")
    with pytest.raises(ValueError, match="没有"):
        reference_text(b"  ", "empty.txt")


def test_reference_xlsx_omits_blank_rows_and_trailing_cells():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Fields"
    sheet.cell(row=3, column=1, value="Invoice Date")
    sheet.cell(row=3, column=2, value="Date")
    sheet.cell(row=20, column=5, value=None)
    out = io.BytesIO()
    book.save(out)
    parsed = json.loads(reference_text(out.getvalue(), "reference.xlsx"))
    assert parsed == [{"sheet": "Fields", "rows": [{"row": 3, "values": ["Invoice Date", "Date"]}]}]


def test_reference_analysis_job_end_to_end(tmp_path, monkeypatch):
    async def fake_completion(c, key, instructions, images, reference_text=""):
        assert "REFERENCE_ONLY" in reference_text
        assert "REFERENCE_ONLY" not in instructions
        assert images and "单据类型" in instructions
        return json.dumps({"profile": {"accepted_documents": "发票"}, "fields": [
            {"table_name": "AI_Document", "name": "InvoiceDate", "data_type": "String"}]})
    monkeypatch.setattr("datara.app.completion_retrying", fake_completion)
    with TestClient(create_app(tmp_path)) as client:
        ref = client.post("/api/references", files={"file": ("rules.txt", b"REFERENCE_ONLY")}).json()
        assert "text" not in ref
        assert client.get("/api/references/" + ref["id"]).json()["name"] == "rules.txt"
        data = io.BytesIO(); Image.new("RGB", (30, 30), "white").save(data, "PNG")
        sample = client.post("/api/samples", files={"file": ("page.png", data.getvalue())}).json()
        client.post("/api/settings", json={"base_url": "https://example.test/v1", "model": "test", "api_key": "test-key"})
        p = new_profile(); p.reference_ids = [ref["id"]]
        job = client.post("/api/jobs", json={"profile": p.model_dump(), "sample_id": sample["id"], "kind": "draft"}).json()
        for _ in range(100):
            result = client.get("/api/jobs/" + job["id"]).json()
            if result["status"] != "running":
                break
            time.sleep(.01)
        assert result["status"] == "completed"
        assert result["profile_suggestion"]["accepted_documents"] == "发票"
        assert result["suggestions"][0]["field"]["data_type"] == "Date"
        assert result["suggestions"][0]["field"]["head_display"] == 1
        p.tables[0].fields = [FieldDef(**result["suggestions"][0]["field"])] + p.tables[0].fields
        saved = client.post("/api/profiles/save", json=p.model_dump()).json()
        assert client.post("/api/export", json=saved).content.startswith(b"PK")
