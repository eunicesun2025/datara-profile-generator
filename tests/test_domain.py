import io
import json
import re
import zipfile

import pytest
from openpyxl import load_workbook

from datara.domain import (FieldDef, TableDef, effective_sql, model_json, new_profile,
                           normalize, strict_json, validate_profile, validate_result)
from datara.generators import export_zip, json_structure, mapping_rows, prompt, sql


@pytest.fixture
def profile():
    p = new_profile("Test", "AI_Invoice_Head")
    p.accepted_documents = "供应商发票"
    p.tables[0].fields += [FieldDef(name="invoice_date", data_type="Date", is_required=True),
                           FieldDef(name="invoice_amount", data_type="Decimal"),
                           FieldDef(name="company_code", source="System", population="Datara 查询"),
                           FieldDef(name="private_note", source="Manual", extraction="MUST_NOT_LEAK")]
    p.tables.append(TableDef(name="AI_POGR", role="detail", parent_table_id=p.tables[0].id,
                            fields=[FieldDef(name="po_number"), FieldDef(name="gr_number")]))
    normalize(p)
    return p


def test_artifacts_have_identical_field_sets_and_no_system_leak(profile):
    archive = zipfile.ZipFile(io.BytesIO(export_zip(profile)))
    assert set(archive.namelist()) == {"field_mapping.xlsx", "create_tables.sql", "extraction_prompt.txt", "output_structure.json"}
    ws = load_workbook(io.BytesIO(archive.read("field_mapping.xlsx"))).active
    rows = list(ws.values)[1:]
    ddl = archive.read("create_tables.sql").decode()
    for table in profile.tables:
        mapped = {r[3] for r in rows if r[0] == table.name}
        block = ddl.split(f"CREATE TABLE [dbo].[{table.name}] (")[1].split("\n);")[0]
        columns = set(re.findall(r"^    \[(\w+)\]", block, re.M))
        assert columns == mapped == {f.name for f in table.fields}
    ai_json = json.loads(archive.read("output_structure.json"))
    assert set(ai_json["AI_Invoice_Head"]) == {"invoice_date", "invoice_amount"}
    assert set(ai_json["AI_POGR"][0]) == {"po_number", "gr_number"}
    text = archive.read("extraction_prompt.txt").decode()
    assert json.dumps(ai_json, ensure_ascii=False, indent=2) in text
    for forbidden in ["file_id", "head_id", "modified_by", "company_code", "private_note", "MUST_NOT_LEAK"]:
        assert forbidden not in text


def test_system_schema_and_nullable_business_fields(profile):
    ddl = sql(profile)
    assert ddl.count("IDENTITY(1,1)") == 2
    assert ddl.count("DEFAULT (SYSUTCDATETIME())") == 2
    assert ddl.count("[modified_by]") == 1
    assert "[invoice_date] date NULL" in ddl
    assert "[invoice_amount] decimal(18,2) NULL" in ddl
    assert "REFERENCES [dbo].[AI_Invoice_Head] ([id]) ON DELETE CASCADE" in ddl
    assert "DROP" not in ddl and "USE" not in ddl


def test_edit_source_changes_every_projection(profile):
    profile.tables[0].fields[-4].source = "System"
    assert "invoice_date" not in prompt(profile)
    assert "[invoice_date] date NULL" in sql(profile)
    assert any(r[3] == "invoice_date" and r[8] == "System" for r in mapping_rows(profile))


def test_null_empty_arrays_and_not_matched(profile):
    value = {"AI_Invoice_Head": {"invoice_date": None, "invoice_amount": None}, "AI_POGR": []}
    result = validate_result(profile, value)
    assert result["status"] == "valid" and len(result["warnings"]) == 1
    assert validate_result(profile, {})["status"] == "not_matched"
    del value["AI_Invoice_Head"]["invoice_date"]
    assert validate_result(profile, value)["errors"]


@pytest.mark.parametrize("date", ["20260230", "2026-05-22", "20261301", "", 20260522])
def test_invalid_dates(profile, date):
    value = {"AI_Invoice_Head": {"invoice_date": date, "invoice_amount": 1}, "AI_POGR": []}
    assert validate_result(profile, value)["status"] == "invalid"


def test_strict_types_extra_keys_and_decimal_precision(profile):
    good = {"AI_Invoice_Head": {"invoice_date": "20260522", "invoice_amount": 32950.0},
            "AI_POGR": [{"po_number": "09300722823", "gr_number": None}]}
    assert validate_result(profile, good)["status"] == "valid"
    good["AI_POGR"][0]["po_number"] = 9300722823
    good["AI_POGR"][0]["head_id"] = 1
    good["AI_Invoice_Head"]["invoice_amount"] = 1.123
    assert len(validate_result(profile, good)["errors"]) == 3


@pytest.mark.parametrize("name", ["Invoice_Date", "invoiceDate", "invoice__date", "invoice_date;DROP", "123_date"])
def test_snake_case_validation(profile, name):
    profile.tables[0].fields[-4].name = name
    assert validate_profile(profile)["errors"]


def test_unsafe_sql_and_bad_structure_are_rejected(profile):
    profile.tables[0].fields[-4].sql_type = "date); DROP TABLE users;--"
    profile.tables[1].parent_table_id = profile.tables[1].id
    with pytest.raises(ValueError):
        export_zip(profile)


def test_reserved_system_fields_cannot_be_reclassified(profile):
    profile.tables[1].fields[-1].source = "AI"  # head_id
    assert validate_profile(profile)["errors"]
    normalize(profile, repair_system=True)
    assert not validate_profile(profile)["errors"]


def test_xlsx_preserves_formula_like_samples_as_literal(profile):
    profile.tables[0].fields[-1].sample_data = '=HYPERLINK("https://example.test")'
    z = zipfile.ZipFile(io.BytesIO(export_zip(profile)))
    ws = load_workbook(io.BytesIO(z.read("field_mapping.xlsx"))).active
    cell = next(row[4] for row in ws if row[3].value == "private_note")
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_duplicate_keys_non_json_and_nonfinite_are_rejected():
    for value in ['{"a":1,"a":2}', '{"a":NaN}', '```json\n{}\n```']:
        with pytest.raises(ValueError):
            strict_json(value)


def test_model_json_accepts_only_a_full_response_code_fence():
    """A fenced model output used to be scored as empty despite holding the right value."""
    assert model_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert model_json('```\n{"a": 1}\n```') == {"a": 1}
    assert model_json('{"a": 1}') == {"a": 1}
    # The fence must enclose the entire response, and strictness survives inside it.
    for value in ['说明 ```json\n{"a":1}\n```', '```json\n{"a":1}\n``` 说明',
                  '```json\n{"a":1,"a":2}\n```', '```json\nNaN\n```', '```json\n{"a":1}']:
        with pytest.raises(ValueError):
            model_json(value)


def test_conflicting_missing_value_rule_is_rejected(profile):
    profile.document_rules = "如果找不到税额，默认返回0。"
    assert any("null" in e for e in validate_profile(profile)["errors"])
