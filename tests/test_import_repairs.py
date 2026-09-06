import io
import zipfile

import pytest
from openpyxl import Workbook
from fastapi.testclient import TestClient

from datara.app import create_app
from datara.domain import validate_profile
from datara.importer import inspect_workbook, parse_fields


def legacy_mapping(parent='', conflicting_parent=False):
    wb = Workbook()
    ws = wb.active
    ws.append(['TableName', 'TableLevel', 'ForeignKeyField', 'ColumnName', 'DataType', 'Source'])
    ws.append(['Header', 1, None, 'date', 'Date', 'AI'])
    ws.append(['Lines', 2, parent, 'amount', 'Decimal', 'AI'])
    ws.append(['Header', 2, None, 'validity_check', 'String', 'Manual'])
    if conflicting_parent:
        ws.append(['Lines', 2, 'Other', 'note', 'String', 'Manual'])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_repairs_require_opt_in_and_preserve_field_table():
    raw = legacy_mapping()
    with pytest.raises(ValueError, match='第 4 行'):
        parse_fields(raw, 'Sheet')
    inspection = inspect_workbook(raw)['sheets'][0]
    assert any('validity_check' in n for n in inspection['repair_notes'])
    p, notes = parse_fields(raw, 'Sheet', repair_structure=True)
    assert not validate_profile(p)['errors']
    assert p.tables[1].parent_table_id == p.tables[0].id
    assert any(f.name == 'validity_check' for f in p.tables[0].fields)
    assert '导入修复记录' in p.description
    assert any('唯一主表' in n for n in notes)


def test_repairs_do_not_overwrite_explicit_wrong_parent():
    with pytest.raises(ValueError, match='ForeignKeyField'):
        parse_fields(legacy_mapping('Unknown'), 'Sheet', repair_structure=True)
    with pytest.raises(ValueError, match='不一致'):
        parse_fields(legacy_mapping('Header', True), 'Sheet', repair_structure=True)


def test_import_rename_save_export_roundtrip(tmp_path):
    with TestClient(create_app(tmp_path)) as c:
        r = c.post('/api/import/inspect', files={'file': ('legacy.xlsx', legacy_mapping())})
        assert r.status_code == 200
        request = {'import_id': r.json()['import_id'], 'sheet': 'Sheet'}
        assert c.post('/api/import/apply', json=request).status_code == 400
        p = c.post('/api/import/apply', json={**request, 'repair_structure': True}).json()['profile']
        p['tables'][0]['name'] = 'Renamed_Header'
        p = c.post('/api/profiles/save', json=p).json()
        assert c.get('/api/profiles/'+p['id']).json()['tables'][0]['name'] == 'Renamed_Header'
        z = zipfile.ZipFile(io.BytesIO(c.post('/api/export', json=p).content))
        assert 'REFERENCES [dbo].[Renamed_Header]' in z.read('create_tables.sql').decode()
        assert 'Renamed_Header' in z.read('extraction_prompt.txt').decode()
        assert 'validity_check' not in z.read('extraction_prompt.txt').decode()
