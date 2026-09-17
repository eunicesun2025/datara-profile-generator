import io
import json
import logging
import os
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from datara.app import (configure_logging, create_app, load_runtime_environment,
                        resolve_log_level)
from datara.domain import FieldDef, new_profile
from datara.importer import inspect_workbook, parse_fields


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as c:
        yield c


@pytest.fixture
def restore_root_logger():
    """Snapshot root logger state so level tests cannot leak into the rest of the suite."""
    root = logging.getLogger()
    level, handlers = root.level, list(root.handlers)
    yield root
    root.setLevel(level)
    root.handlers[:] = handlers


def test_project_env_loads_once_without_overriding_environment(tmp_path, monkeypatch):
    (tmp_path / ".env").write_text(
        "DATARA_API_KEY=local-project-key\nDATARA_DATA_DIR=local-data\n",
        encoding="utf-8",
    )
    monkeypatch.delenv("DATARA_API_KEY", raising=False)
    monkeypatch.setenv("DATARA_DATA_DIR", "deployment-data")

    load_runtime_environment(tmp_path)

    assert os.environ["DATARA_API_KEY"] == "local-project-key"
    assert os.environ["DATARA_DATA_DIR"] == "deployment-data"


@pytest.mark.parametrize("raw,expected", [
    ("DEBUG", "DEBUG"),
    ("info", "INFO"),          # 大小写不敏感，.env 里手写小写也应生效
    ("  Warning  ", "WARNING"),  # 容忍空白
    ("ERROR", "ERROR"),
    ("CRITICAL", "CRITICAL"),
    (None, "WARNING"),         # 未设置 → 默认，保持改动前的行为
    ("", "WARNING"),           # .env 里留空 → 默认
    ("verbose", "WARNING"),    # 拼错 → 默认，而不是启动崩溃
])
def test_resolve_log_level_accepts_known_names_and_falls_back(raw, expected):
    assert resolve_log_level(raw) == expected


def test_configure_logging_keeps_warning_by_default(monkeypatch, restore_root_logger):
    """默认必须是 WARNING：否则本次改动会让所有部署突然开始刷屏。"""
    monkeypatch.delenv("DATARA_LOG_LEVEL", raising=False)
    assert configure_logging() == "WARNING"
    assert restore_root_logger.level == logging.WARNING


def test_configure_logging_enables_provider_diagnostics_at_info(monkeypatch, restore_root_logger):
    """DATARA_LOG_LEVEL=INFO 必须让 datara.provider 的 INFO 记录真正可见。

    回归防护：uvicorn 的 --log-level 只对 uvicorn.error/access/asgi 调 setLevel
    (uvicorn/config.py:413-420)，从不配置 root。缺了 configure_logging() 时这些计时行会
    落到 logging.lastResort(WARNING) 被丢弃 —— 这正是 600 秒代理阻塞当初无法定位的原因。
    """
    monkeypatch.setenv("DATARA_LOG_LEVEL", "INFO")
    assert configure_logging() == "INFO"
    assert logging.getLogger("datara.provider").getEffectiveLevel() == logging.INFO

    seen = []
    handler = logging.Handler()
    handler.emit = seen.append
    restore_root_logger.addHandler(handler)
    try:
        logging.getLogger("datara.provider").info("model_request_started")
        logging.getLogger("datara.provider").debug("below-threshold")
    finally:
        restore_root_logger.removeHandler(handler)

    # 只断言 INFO 那条：DEBUG 仍应低于阈值，证明级别是 INFO 而不是被放成 DEBUG。
    assert [r.getMessage() for r in seen] == ["model_request_started"]


def test_configure_logging_does_not_duplicate_root_handlers(monkeypatch, restore_root_logger):
    """重复调用不得叠加 handler，否则每行日志会被打印多次。"""
    monkeypatch.setenv("DATARA_LOG_LEVEL", "INFO")
    configure_logging()
    first = len(restore_root_logger.handlers)
    configure_logging()
    configure_logging()
    assert len(restore_root_logger.handlers) == first


def test_configure_logging_warns_about_an_unusable_level(monkeypatch, restore_root_logger, caplog):
    """拼错的级别要给出可见提示，而不是静默按 WARNING 运行让人以为设置生效了。"""
    monkeypatch.setenv("DATARA_LOG_LEVEL", "verbose")
    with caplog.at_level(logging.WARNING, logger="datara.app"):
        assert configure_logging() == "WARNING"
    assert any("DATARA_LOG_LEVEL" in r.getMessage() and "verbose" in r.getMessage()
               for r in caplog.records)


def test_save_reload_conflict_and_export(client):
    p = client.post("/api/demo/invoice").json()
    saved = client.post("/api/profiles/save", json=p).json()
    assert saved["revision"] == 1
    loaded = client.get("/api/profiles/" + saved["id"]).json()
    assert loaded == saved
    assert client.post("/api/profiles/save", json=p).status_code == 409
    preview = client.post("/api/preview", json=saved).json()
    assert preview["errors"] == []
    assert client.post("/api/export", json=saved).content.startswith(b"PK")
    saved["tables"][0]["fields"][0]["name"] = "renamed_field"
    assert client.post("/api/export", json=saved).status_code == 409


def test_settings_never_persist_or_return_api_key(client, tmp_path):
    data = {"base_url": "https://example.test/v1", "model": "test-vision", "api_key": "test-secret-token"}
    response = client.post("/api/settings", json=data)
    assert response.status_code == 200
    assert "test-secret-token" not in response.text
    assert "test-secret-token" not in (client.app.state.store.root / "connection.json").read_text()
    assert client.get("/api/settings").json()["has_key"] is True
    assert client.post("/api/settings", json={**data, "clear_key": True}).json()["has_key"] is False


def test_reject_cross_origin_and_path_traversal(client):
    assert client.post("/api/profiles/new", headers={"Origin": "https://evil.test"}).status_code == 403
    assert client.get("/api/profiles/not.valid.id").status_code == 400
    assert client.get("/api/health", headers={"Host": "evil.test"}).status_code == 403


def image_bytes():
    out = io.BytesIO()
    Image.new("RGB", (120, 100), "white").save(out, "PNG")
    return out.getvalue()


def test_image_upload_preview_and_invalid_file(client):
    response = client.post("/api/samples", files={"file": ("sample.png", image_bytes(), "image/png")})
    assert response.status_code == 200
    sample = response.json()
    assert sample["pages"] == 1
    assert client.get(f"/api/samples/{sample['id']}/pages/1").headers["content-type"] == "image/jpeg"
    assert client.get(f"/api/samples/{sample['id']}/pages/2").status_code == 404
    assert client.post("/api/samples", files={"file": ("bad.pdf", b"not a pdf", "application/pdf")}).status_code == 400


def spreadsheet():
    w = Workbook(); s = w.active
    s.title = "Mapping"
    s.append(["TableName", "TableLevel", "ForeignKeyField", "ColumnName", "SampleData", "DataType", "ChoiceValues", "IsRequired", "Source", "HeadDisplay", "FieldOrder"])
    s.append(["AI_Invoice_Head", 1, None, "invoice_number", "0012", "String", None, "true", "AI", 1, 2])
    s.append(["AI_Invoice_Head", 1, None, "company_code", None, "String", None, None, "AI", None, 2])
    s.append(["AI_POGR", 2, "AI_Invoice_Head", "po_number", "09300000000", "String", None, None, "AI", None, 1])
    s.append(["AI_POGR", 2, "AI_Invoice_Head", "head_id", None, "Integer", None, None, "AI", None, 2])
    out=io.BytesIO(); w.save(out); return out.getvalue()


def test_mapping_import_repairs_legacy_and_preserves_identifiers(client):
    info=client.post("/api/import/inspect", files={"file": ("mapping.xlsx", spreadsheet())}).json()
    result=client.post("/api/import/apply", json={"import_id":info["import_id"],"sheet":"Mapping"}).json()
    p=result["profile"]
    assert result["notes"]
    assert p["tables"][0]["fields"][0]["sample_data"] == "0012"
    assert p["tables"][0]["fields"][1]["source"] == "System"
    assert next(f for f in p["tables"][1]["fields"] if f["name"] == "head_id")["source"] == "System"
    assert client.post("/api/preview", json=p).json()["errors"] == []


def test_plain_excel_column_selection():
    w=Workbook();s=w.active;s.title="业务清单";s.append(["名称","说明"]);s.append(["account_no","账号"])
    out=io.BytesIO();w.save(out)
    p, notes=parse_fields(out.getvalue(),"业务清单",{"name":"名称","description":"说明"})
    assert p.tables[0].fields[0].description == "账号"
    assert not p.tables[0].fields[0].reviewed


def test_unknown_mapping_source_is_not_marked_reviewed():
    from openpyxl import load_workbook
    w = load_workbook(io.BytesIO(spreadsheet()))
    w.active["I2"] = "Unknown"
    out = io.BytesIO(); w.save(out)
    p, notes = parse_fields(out.getvalue(), "Mapping")
    assert p.tables[0].fields[0].reviewed is False
    assert any("请审核" in note for note in notes)


def test_async_vision_job_records_validation_without_credentials(client, monkeypatch):
    async def fake_completion(c,key,instructions,images,reference_text=""):
        assert key == "test-only-key"
        assert images and images[0].exists()
        assert "file_id" not in instructions
        return '{"AI_Document":{"amount":null}}'
    monkeypatch.setattr("datara.app.completion_retrying",fake_completion)
    client.post("/api/settings",json={"base_url":"https://example.test/v1","model":"test","api_key":"test-only-key"})
    sample=client.post("/api/samples",files={"file":("a.png",image_bytes())}).json()
    p=new_profile();p.tables[0].fields.append(FieldDef(name="amount",data_type="Decimal",is_required=True))
    job=client.post("/api/jobs",json={"profile":p.model_dump(),"sample_id":sample["id"],"kind":"extract"}).json()
    for _ in range(30):
        result=client.get("/api/jobs/"+job["id"]).json()
        if result["status"] != "running": break
        time.sleep(.01)
    assert result["status"] == "completed"
    assert result["validation"]["status"] == "valid"
    assert result["validation"]["warnings"]
    assert "test-only-key" not in client.app.state.store.path("tests",job["id"]).read_text()


def test_extract_job_survives_a_transient_connection_failure(client, monkeypatch):
    """Wiring check: run_job must call the model through the retrying wrapper.

    A single 15 second connect timeout from a corporate proxy used to fail the whole
    测试提取 run, even though the message shown to the user said it was retryable.
    """
    calls={"n":0}
    async def flaky(c,key,instructions,images,reference_text=""):
        calls["n"]+=1
        if calls["n"]==1:
            raise ValueError("无法连接模型端点：建立连接超过 15 秒连接超时（企业代理握手缓慢或网络抖动），可重试")
        return '{"AI_Document":{"amount":null}}'
    monkeypatch.setattr("datara.provider.completion",flaky)
    client.post("/api/settings",json={"base_url":"https://example.test/v1","model":"test","api_key":"test-only-key"})
    sample=client.post("/api/samples",files={"file":("a.png",image_bytes())}).json()
    p=new_profile();p.tables[0].fields.append(FieldDef(name="amount",data_type="Decimal",is_required=True))
    job=client.post("/api/jobs",json={"profile":p.model_dump(),"sample_id":sample["id"],"kind":"extract"}).json()
    result={}
    for _ in range(120):
        result=client.get("/api/jobs/"+job["id"]).json()
        if result["status"] != "running": break
        time.sleep(.05)
    assert calls["n"] == 2
    assert result["status"] == "completed"


def test_extract_job_parses_a_fenced_model_response(client, monkeypatch):
    """A model response wrapped in a Markdown fence must not fail 测试提取.

    Regression: fenced responses parsed to {}, the page showed
    'Expecting value: line 1 column 1 (char 0)' and an empty result even though the
    extraction itself had succeeded.
    """
    async def fenced(c,key,instructions,images,reference_text=""):
        return '```json\n{"AI_Document":{"amount":null}}\n```'
    monkeypatch.setattr("datara.app.completion_retrying",fenced)
    client.post("/api/settings",json={"base_url":"https://example.test/v1","model":"test","api_key":"test-only-key"})
    sample=client.post("/api/samples",files={"file":("a.png",image_bytes())}).json()
    p=new_profile();p.tables[0].fields.append(FieldDef(name="amount",data_type="Decimal",is_required=True))
    job=client.post("/api/jobs",json={"profile":p.model_dump(),"sample_id":sample["id"],"kind":"extract"}).json()
    result={}
    for _ in range(30):
        result=client.get("/api/jobs/"+job["id"]).json()
        if result["status"] != "running": break
        time.sleep(.01)
    assert result["status"] == "completed"
    assert result["result"] == {"AI_Document": {"amount": None}}
    assert result["validation"]["status"] == "valid"
