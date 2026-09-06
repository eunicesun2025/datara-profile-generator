import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from .domain import FieldDef, Profile, TableDef, normalize


def workbook(content: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024 or len(z.infolist()) > 3000:
                raise ValueError("Excel 解压后过大，请缩小字段清单")
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError) as e:
        raise ValueError("请上传有效的 .xlsx 文件") from e


def inspect_workbook(content: bytes) -> dict:
    wb = workbook(content)
    try:
        sheets = []
        for ws in wb:
            if ws.max_row > 10000 or ws.max_column > 100:
                raise ValueError("字段清单最多支持 10,000 行和 100 列")
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True))
            sheets.append({"name": ws.title, "headers": [str(v or "") for v in (rows[0] if rows else [])],
                           "preview": [[scalar(v) for v in r] for r in rows[1:]], "rows": ws.max_row})
        for sheet in sheets:
            if all(k in sheet['headers'] for k in ['TableName', 'TableLevel', 'ColumnName', 'Source', 'DataType']):
                try:
                    _, notes = parse_fields(content, sheet['name'], repair_structure=True)
                    sheet['repair_notes'] = notes
                except ValueError as e:
                    sheet['import_error'] = str(e)
        return {"sheets": sheets}
    finally:
        wb.close()


def scalar(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y%m%d")
    return v if v is None or isinstance(v, (str, int, float, bool)) else str(v)


def as_bool(v):
    return str(v).strip().lower() in {"true", "1", "yes", "是", "y"}


def parse_fields(content: bytes, sheet: str, columns: dict | None = None, repair_structure: bool = False) -> tuple[Profile, list[str]]:
    wb = workbook(content)
    notes = []
    try:
        if sheet not in wb.sheetnames:
            raise ValueError("所选工作表不存在")
        ws = wb[sheet]
        if ws.max_row > 10000 or ws.max_column > 100:
            raise ValueError("字段清单过大")
        rows = ws.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(rows, [])]
        full = all(k in headers for k in ["TableName", "TableLevel", "ColumnName", "Source", "DataType"])
        if not full and (not columns or "name" not in columns):
            raise ValueError("普通字段清单请选择字段名称对应列")
        p = Profile(name="导入的 Profile")
        tables, parents = {}, {}
        for line, values in enumerate(rows, 2):
            def get(column, default=None):
                if column not in headers:
                    return default
                index = headers.index(column)
                return values[index] if index < len(values) and values[index] is not None else default

            name = get("ColumnName") if full else get(columns["name"])
            if name is None or str(name).strip() == "":
                continue
            name = str(name).strip()
            tname = str(get("TableName", "AI_Document")) if full else "AI_Document"
            level = str(get("TableLevel", 1)) if full else "1"
            if level not in {"1", "2"}:
                raise ValueError(f"第 {line} 行：只支持 TableLevel 1 或 2")
            role = "head" if level == "1" else "detail"
            parent = str(get("ForeignKeyField", "")) if full else ""
            if tname not in tables:
                tables[tname] = TableDef(name=tname, role=role)
                parents[tname] = parent
            elif tables[tname].role != role or parents[tname] != parent:
                if not repair_structure or (parent and parents[tname] and parents[tname] != parent):
                    raise ValueError(f"第 {line} 行：同一表的层级或父表不一致")
                if tables[tname].role != role:
                    notes.append(f"第 {line} 行 {name}：保留表名 {tname}，层级按该表首次出现的定义改为 {'1（主表）' if tables[tname].role == 'head' else '2（子表）'}；请审核字段归属")
                if parent and not parents[tname]:
                    parents[tname] = parent
            dtype = str(get("DataType", "") if full else get(columns.get("type"), ""))
            type_map = {"string": "String", "integer": "Integer", "decimal": "Decimal", "date": "Date",
                        "boolean": "Boolean", "choice": "Choice", "文本": "String", "金额": "Decimal", "日期": "Date"}
            recognized = dtype.lower() in type_map
            dtype = type_map.get(dtype.lower(), "String")
            source = str(get("Source", "") if full else get(columns.get("source"), ""))
            source_known = source.lower() in {"ai", "system", "manual"}
            source = {"ai": "AI", "system": "System", "manual": "Manual"}.get(source.lower(), "AI")
            reviewed = full and recognized and source_known
            if not reviewed:
                notes.append(f"第 {line} 行 {name}：来源/类型请审核，未提供时暂以 AI / String 建议")
            if name in {"company_code", "current_date"} or (name == "item" and tname == "AI_Invoice_Detail"):
                if source != "System":
                    notes.append(f"{tname}.{name}：按确认规则改为 System，由 Datara 填充")
                source = "System"
            display = get("HeadDisplay") if full else None
            if display not in (None, ""):
                try:
                    display = int(display)
                except (ValueError, TypeError):
                    raise ValueError(f"第 {line} 行：HeadDisplay 必须为整数")
            field = FieldDef(name=name, source=source, data_type=dtype, reviewed=reviewed,
                             description=str(get(columns.get("description"), "")) if not full else "",
                             sample_data=scalar(get("SampleData")) if full else None,
                             is_required=as_bool(get("IsRequired", False)) if full else False,
                             choice_values=[s.strip() for s in str(get("ChoiceValues", "")).replace("，", ",").split(",") if s.strip()] if full else [],
                             head_display=display,
                             population="由 Datara 查询或生成" if source == "System" else "")
            tables[tname].fields.append(field)
            if len(tables[tname].fields) > 190:
                raise ValueError("每张表最多导入 190 个字段，为必要系统字段保留空间")
        if not tables:
            raise ValueError("没有找到可导入的字段")
        for table in tables.values():
            if table.role == "detail":
                heads = [t for t in tables.values() if t.role == 'head']
                if not parents[table.name] and repair_structure and len(heads) == 1:
                    parents[table.name] = heads[0].name
                    notes.append(f"{table.name}：空白 ForeignKeyField 补为唯一主表 {heads[0].name}（此列填写主表名）")
                parent = tables.get(parents[table.name])
                if parent is None or parent.role != "head":
                    raise ValueError(f"{table.name}：ForeignKeyField 应填写已存在的主表名称")
                table.parent_table_id = parent.id
        p.tables = list(tables.values())
        if repair_structure and notes:
            p.description = ('导入修复记录（字段归属待审核）：\n' + '\n'.join(notes))[:2000]
            for table in p.tables:
                for field in table.fields:
                    field.reviewed = False
        notes += normalize(p, repair_system=True)
        notes.append("FieldOrder 已按工作表行次序逐表连续编号；SQL 类型使用已确认默认值。请审核后保存。")
        return p, notes
    finally:
        wb.close()
